# -*- coding: utf-8 -*-
"""路演 PPT 构建器 v2（浅色主题 · 15 页）：python submission/build_pptx.py
内容与 05_路演PPT_初稿.md 一致，证据只用真实验证过的数字。"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── 浅色主题色板（与 design_tokens.py 一致）──
BG = RGBColor(0xF7, 0xF8, 0xFA)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE9, 0xE7, 0xE3)
FG = RGBColor(0x1A, 0x1A, 0x2E)
MUTED = RGBColor(0x6B, 0x72, 0x80)
MUTED_SOFT = RGBColor(0x9C, 0xA3, 0xAF)
PRIMARY = RGBColor(0x5B, 0x63, 0xE8)
PRIMARY_SOFT = RGBColor(0xEE, 0xF0, 0xFF)
EXPERT = RGBColor(0x4F, 0x7D, 0xFF)
EXPERT_SOFT = RGBColor(0xEE, 0xF3, 0xFF)
RED = RGBColor(0xE8, 0x5B, 0x5B)
RED_SOFT = RGBColor(0xFD, 0xEC, 0xEC)
DECISION = RGBColor(0x7A, 0x65, 0xD8)
DECISION_SOFT = RGBColor(0xF2, 0xEE, 0xFB)
SUCCESS = RGBColor(0x16, 0xA3, 0x4A)
WARNING = RGBColor(0xE8, 0x87, 0x0D)
FONT = "微软雅黑"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

HERE = Path(__file__).resolve().parent


def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def card(s, x, y, w, h, fill=SURFACE, line=BORDER, line_w=1.0, round_=True):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp_after=6):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(sp_after)
        for (txt, size, color, bold) in para:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.name = FONT
            r.font._rPr.set(qn("w:eastAsia"), FONT)
    return tb


def eyebrow(s, txt, x=0.7, y=0.55, color=PRIMARY):
    text(s, x, y, 12, 0.35, [[(txt, 11, color, True)]])


def title(s, txt, x=0.7, y=0.95, size=28, color=FG):
    text(s, x, y, 12, 0.7, [[(txt, size, color, True)]])


def footer(s, n):
    text(s, 0.7, 7.02, 9, 0.35, [[("创想∞ AI创业委员会", 10, MUTED_SOFT, False)]])
    text(s, 12.2, 7.02, 0.8, 0.35, [[(str(n), 10, MUTED_SOFT, False)]], align=PP_ALIGN.RIGHT)


def chip(s, txt, x, y, w, fill, color):
    card(s, x, y, w, 0.36, fill=fill, line=fill)
    text(s, x, y + 0.04, w, 0.3, [[(txt, 11, color, True)]], align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════
# P1 封面
# ═══════════════════════════════════════════════
s = slide()
# 顶部主色条
card(s, 0, 0, 13.333, 0.16, fill=PRIMARY, line=PRIMARY, round_=False)
# 主标题
text(s, 0.9, 2.3, 11.5, 1.4, [[("创想∞ AI创业委员会", 52, FG, True)]])
text(s, 0.9, 3.45, 11.5, 0.8, [[("让 AI 先质疑你的创业想法", 24, PRIMARY, False)]])
text(s, 0.9, 4.4, 11.5, 1.2,
     [[("不是帮你把商业计划写得更漂亮，而是先让它经受一场委员会的挑战。", 15, MUTED, False)]])
# 底部标签
labels = [("12 Agent", EXPERT_SOFT, EXPERT), ("三委员会", RED_SOFT, RED),
          ("5层DAG", DECISION_SOFT, DECISION), ("真实全链路运行", PRIMARY_SOFT, PRIMARY)]
xx = 0.9
for t, f, c in labels:
    chip(s, t, xx, 6.3, 1.8, f, c)
    xx += 2.0

# ═══════════════════════════════════════════════
# P2 一句话定位
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "POSITIONING")
title(s, "这是一个普通的 AI 写作工具吗？")
# 大答句
text(s, 0.7, 2.2, 12, 1.0, [[("不——这是一个会“唱反调”的 AI 创业委员会", 32, PRIMARY, True)]])
# 三行支撑
points = [
    ("不默认认可你的假设", "主动寻找它为什么可能失败，而非顺着你扩写"),
    ("12 位角色分工冲突的 Agent", "不是一个模型一口气写完，而是委员会接力审议"),
    ("输出创业项目诊断报告", "可追溯的决策链与 12 份报告，而非一篇漂亮文案"),
]
for i, (h, b) in enumerate(points):
    y = 3.6 + i * 1.0
    card(s, 0.7, y, 11.9, 0.82, fill=SURFACE)
    text(s, 1.0, y + 0.12, 3.2, 0.6, [[(h, 16, FG, True)]])
    text(s, 4.2, y + 0.18, 8.2, 0.55, [[(b, 14, MUTED, False)]])
footer(s, 2)

# ═══════════════════════════════════════════════
# P3 用户痛点
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "PAIN POINTS")
title(s, "早期创业最缺的不是鼓励，而是敢说“不”的人")
pains = [
    ("过度相信自己判断", "需求没验证就开干，把“我觉得”当成“用户需要”"),
    ("忽略真实需求与付费方", "使用者和买单的不是同一个人，付费证据为零"),
    ("高估市场、忽视竞品与壁垒", "“我们没有竞品”是最危险的信号"),
    ("忽视商业模式、财务与增长", "没有冷启动路径，上线即静默"),
]
for i, (h, b) in enumerate(pains):
    x = 0.7 + (i % 2) * 6.1
    y = 2.1 + (i // 2) * 2.0
    card(s, x, y, 5.85, 1.7, fill=SURFACE)
    card(s, x, y, 5.85, 0.1, fill=RED, line=RED, round_=False)
    text(s, x + 0.3, y + 0.28, 5.3, 0.5, [[(h, 17, FG, True)]])
    text(s, x + 0.3, y + 0.88, 5.3, 0.7, [[(b, 13, MUTED, False)]])
text(s, 0.7, 6.3, 12, 0.5,
     [[("传统 AI 往往顺着用户继续完善，而不是主动反驳。", 14, RED, True)]])
footer(s, 3)

# ═══════════════════════════════════════════════
# P4 项目解决方案
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "SOLUTION")
title(s, "从“输入想法”到“拿到诊断报告”的完整闭环")
steps = [
    ("输入\n创业想法", PRIMARY),
    ("专家委员会\n8 官并行", EXPERT),
    ("红队\n质疑", RED),
    ("风险\n红线把关", RED),
    ("总指挥\n阶段结论", DECISION),
    ("项目评审\n终审", DECISION),
    ("模拟路演\n答辩", DECISION),
    ("创业项目\n诊断报告", PRIMARY),
]
for i, (st_, c) in enumerate(steps):
    x = 0.45 + i * 1.6
    card(s, x, 2.5, 1.42, 1.5, fill=SURFACE, line=c, line_w=1.5)
    parts = st_.split("\n")
    text(s, x + 0.06, 2.62, 1.3, 1.25,
         [[(p, 13 if j == 0 else 11, FG if j == 0 else MUTED, j == 0)] for j, p in enumerate(parts)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sp_after=2)
card(s, 0.45, 4.5, 12.4, 0.6, fill=PRIMARY_SOFT, line=PRIMARY_SOFT)
text(s, 0.65, 4.58, 12, 0.45,
     [[("第一层 8 个 Agent 由依赖图自动调度、同时开工；后四层严格等待前置层完成。", 13, PRIMARY, True)]])
text(s, 0.7, 5.5, 12, 0.8,
     [[("用户只做一件事——提交想法；剩下的是委员会自己跑完。", 15, FG, False)]])
footer(s, 4)

# ═══════════════════════════════════════════════
# P5 12 Agent 三委员会架构
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "ARCHITECTURE")
title(s, "12 个 Agent，三个委员会，角色冲突被刻意设计")
groups = [
    ("专家委员会 · 8 官并行尽调", EXPERT, EXPERT_SOFT,
     ["用户洞察官", "市场分析官", "竞品分析官", "产品设计官",
      "商业模式官", "财务分析官", "增长运营官", "风险审查官"]),
    ("对抗委员会 · 只攻击不安慰", RED, RED_SOFT, ["红队质疑官"]),
    ("决策委员会 · 分权裁决", DECISION, DECISION_SOFT,
     ["创业总指挥", "项目评审官", "路演答辩官"]),
]
x = 0.7
for h, c, soft, agents in groups:
    n = len(agents)
    h_box = max(2.6, 0.7 + n * 0.46)
    card(s, x, 2.0, 3.95, h_box + 0.5, fill=SURFACE)
    card(s, x, 2.0, 3.95, 0.5, fill=soft, line=soft, round_=False)
    text(s, x + 0.25, 2.08, 3.5, 0.4, [[(h, 13, c, True)]])
    text(s, x + 0.25, 2.7, 3.5, h_box,
         [[("• " + a, 13, FG, False)] for a in agents], sp_after=7)
    x += 4.18
text(s, 0.7, 6.3, 12, 0.5,
     [[("专家分析 · 红队攻击 · 风险否决 · 指挥官与评审官分权——一个系统里必须有人唱反调。", 14, MUTED, False)]])
footer(s, 5)

# ═══════════════════════════════════════════════
# P6 核心工作流：5 层 DAG
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "WORKFLOW")
title(s, "不是写死顺序，而是依赖图驱动的 5 层执行 DAG")
layers = [
    ("L0 专家层（8 并行）", "用户 · 市场 · 竞品 · 产品 · 商业模式 · 财务 · 增长 · 风险", EXPERT, EXPERT_SOFT),
    ("L1 红队层", "红队质疑官", RED, RED_SOFT),
    ("L2 总指挥层", "创业总指挥", DECISION, DECISION_SOFT),
    ("L3 评审层", "项目评审官", DECISION, DECISION_SOFT),
    ("L4 路演层", "路演答辩官", DECISION, DECISION_SOFT),
]
for i, (name, desc, c, soft) in enumerate(layers):
    y = 2.05 + i * 0.92
    card(s, 0.7, y, 12, 0.78, fill=SURFACE)
    card(s, 0.7, y, 0.12, 0.78, fill=c, line=c, round_=False)
    text(s, 1.0, y + 0.1, 2.6, 0.6, [[(name, 14, c, True)]])
    text(s, 3.7, y + 0.16, 8.8, 0.5, [[(desc, 13, MUTED, False)]])
    if i < len(layers) - 1:
        text(s, 0.74, y + 0.74, 0.1, 0.18, [[("▼", 9, c, True)]], align=PP_ALIGN.CENTER)
text(s, 0.7, 6.5, 12, 0.5,
     [[("执行器根据依赖关系自动识别“当前谁可以跑”；硬依赖失败则下游阻断，软依赖缺失则降级。", 13, MUTED, False)]])
footer(s, 6)

# ═══════════════════════════════════════════════
# P7 红队机制（核心卖点）
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "RED TEAM · 核心卖点", color=RED)
title(s, "系统不是只证明想法成立，而是主动寻找它为什么可能失败")
# 五维
dims = ["需求", "付费", "竞争", "增长", "壁垒"]
for i, d in enumerate(dims):
    chip(s, d, 0.7 + i * 1.05, 2.0, 0.9, RED_SOFT, RED)
# 真实攻击原文
attacks = [
    ("【付费假设｜致命】",
     "“口头兴趣 ≠ 采购意愿——预算权、审批流程、立项记录全部为零。”"),
    ("【增长假设｜致命】",
     "“没有冷启动路径，上线即静默；增长目标缺渠道×流量×转化率推导链。”"),
]
for i, (h, q) in enumerate(attacks):
    x = 0.7 + i * 6.1
    card(s, x, 2.6, 5.85, 3.0, fill=SURFACE, line=RED, line_w=1.2)
    card(s, x, 2.6, 5.85, 0.5, fill=RED_SOFT, line=RED_SOFT, round_=False)
    text(s, x + 0.25, 2.68, 5.4, 0.4, [[(h, 14, RED, True)]])
    text(s, x + 0.3, 3.3, 5.35, 2.1, [[(q, 15, FG, False)]])
text(s, 0.7, 5.85, 12, 0.6,
     [[("五维攻击逐维引用项目原文，并给出“待验证证据”——这不是模板话术，是针对该项目生成的攻击。", 13, MUTED, False)]])
footer(s, 7)

# ═══════════════════════════════════════════════
# P8 工程架构
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "ENGINEERING")
title(s, "把多 Agent 系统做成可运行、可恢复、可验证的工程系统")
chain = [
    ("Registry", "管理 Agent 能力与依赖"),
    ("DAG", "编排执行顺序，自动并行"),
    ("Orchestrator", "按图调度执行"),
    ("Validator", "输出结构化校验"),
    ("Critic/Repair", "问题输出返工留证"),
    ("RunStore", "保存运行状态"),
    ("Resume", "断点恢复不重跑"),
    ("Report", "汇总 12 份报告终审"),
]
for i, (name, desc) in enumerate(chain):
    x = 0.7 + (i % 4) * 3.08
    y = 2.1 + (i // 4) * 1.9
    card(s, x, y, 2.85, 1.6, fill=SURFACE)
    text(s, x + 0.2, y + 0.18, 2.5, 0.45, [[(name, 15, PRIMARY, True)]])
    text(s, x + 0.2, y + 0.72, 2.5, 0.7, [[(desc, 12, MUTED, False)]])
    if i % 4 < 3:
        text(s, x + 2.82, y + 0.55, 0.2, 0.4, [[("→", 16, PRIMARY_SOFT and PRIMARY, True)]], align=PP_ALIGN.CENTER)
text(s, 0.7, 6.1, 12, 0.6,
     [[("工程观：LLM 是不受信任的执行者，程序负责最终约束。", 16, PRIMARY, True)]])
footer(s, 8)

# ═══════════════════════════════════════════════
# P9 运行结果（真实量化数据）
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "RESULTS")
title(s, "只展示有证据的数字")
metrics = [
    ("12", "Agent", "8 专家 + 1 红队 + 3 决策", EXPERT),
    ("72", "份知识文档", "支撑 12 Agent 专业判断", PRIMARY),
    ("120", "条 QA 测试", "离线断言全绿", DECISION),
    ("5", "层 DAG", "专家→红队→总指挥→评审→路演", PRIMARY),
    ("12/12", "真实运行完成", "Run 可复查", SUCCESS),
]
for i, (big, label, sub, c) in enumerate(metrics):
    x = 0.7 + i * 2.48
    card(s, x, 2.1, 2.3, 2.6, fill=SURFACE)
    text(s, x, 2.3, 2.3, 0.9, [[(big, 36, c, True)]], align=PP_ALIGN.CENTER)
    text(s, x, 3.25, 2.3, 0.4, [[(label, 13, FG, True)]], align=PP_ALIGN.CENTER)
    text(s, x + 0.1, 3.7, 2.1, 0.8, [[(sub, 11, MUTED, False)]], align=PP_ALIGN.CENTER)
# 单次运行对比
card(s, 0.7, 5.05, 12, 1.1, fill=PRIMARY_SOFT, line=PRIMARY_SOFT)
text(s, 1.0, 5.18, 12, 0.4, [[("单次真实运行对比观察", 13, PRIMARY, True)]])
text(s, 1.0, 5.55, 12, 0.5,
     [[("363.8s  →  267.8s（约 26.4%）", 22, FG, True)]])
text(s, 0.7, 6.25, 12, 0.5,
     [[("* 单次真实运行对比观察，非稳定性能基准；不包装未经统计验证的“准确率/性能指标”。", 11, MUTED, False)]])
footer(s, 9)

# ═══════════════════════════════════════════════
# P10 最终结果展示（真实 Run 结论）
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "FINAL VERDICT · run_ecf846e4ff50")
title(s, "一次真实运行的终审结论")
chain_nodes = [
    ("风险审查官", "风险阻断：是", RED),
    ("创业总指挥", "阶段结论：想法验证", DECISION),
    ("项目评审官", "暂缓进入修订周期\n红线一票否决：是", DECISION),
    ("路演答辩官", "暂不建议路演\n被击穿问题数：24", DECISION),
]
for i, (name, concl, c) in enumerate(chain_nodes):
    x = 0.7 + i * 3.08
    card(s, x, 2.1, 2.85, 2.3, fill=SURFACE)
    card(s, x, 2.1, 2.85, 0.5, fill=DECISION_SOFT if c == DECISION else RED_SOFT,
         line=DECISION_SOFT if c == DECISION else RED_SOFT, round_=False)
    text(s, x + 0.2, 2.18, 2.5, 0.4, [[(name, 13, c, True)]])
    text(s, x + 0.2, 2.8, 2.5, 1.4, [[(concl, 14, FG, True)]], sp_after=4)
# 一致性 + 结算
card(s, 0.7, 4.7, 12, 0.7, fill=SUCCESS and RGBColor(0xEA, 0xFB, 0xEF),
     line=RGBColor(0xBB, 0xE5, 0xC5))
text(s, 1.0, 4.82, 12, 0.5,
     [[("✅ 决策链一致性检查：风险 → 总指挥 → 评审 → 路演 结论链无程序级冲突", 13, SUCCESS, True)]])
card(s, 0.7, 5.55, 12, 0.65, fill=SURFACE)
text(s, 1.0, 5.68, 12, 0.45,
     [[("运行结算：12/12 success ｜ 返工 0 次 ｜ 警告 1 条", 13, FG, False)]])
text(s, 0.7, 6.35, 12, 0.5,
     [[("所有结论来自真实项目报告——系统真的会对你说“不通过”。", 14, RED, True)]])
footer(s, 10)

# ═══════════════════════════════════════════════
# P11 产品体验（UI 截图）
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "PRODUCT EXPERIENCE")
title(s, "从输入想法到拿到结构化诊断结果")
ui_imgs = [
    ("ui_v2/01_home.png", "首页：输入想法"),
    ("ui_v2/02_running.png", "委员会运行 + Agent 状态"),
    ("ui_v2/04_redteam.png", "红队质疑报告"),
    ("ui_v2/03_result.png", "终审报告"),
]
for i, (rel, cap) in enumerate(ui_imgs):
    x = 0.7 + (i % 2) * 6.1
    y = 2.05 + (i // 2) * 2.35
    p = HERE / rel
    if p.exists():
        s.shapes.add_picture(str(p), Inches(x), Inches(y), width=Inches(5.85))
    else:
        card(s, x, y, 5.85, 2.1, fill=SURFACE)
        text(s, x, y + 0.8, 5.85, 0.5, [[("[ 截图占位：" + cap + " ]", 12, MUTED, False)]],
             align=PP_ALIGN.CENTER)
    text(s, x, y + 2.15, 5.85, 0.3, [[(cap, 11, MUTED, True)]], align=PP_ALIGN.CENTER)
text(s, 0.7, 6.5, 12, 0.5,
     [[("用户只做一件事——提交想法；剩下的委员会自己跑完，输出可追溯的诊断结果。", 13, MUTED, False)]])
footer(s, 11)

# ═══════════════════════════════════════════════
# P12 创新与差异化
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "DIFFERENTIATION")
title(s, "不是“用了多 Agent”，而是委员会式决策链")
# 表头
cols = [("维度", 2.6), ("普通 AI 助手", 3.1), ("商业计划生成", 3.1), ("创想∞", 3.5)]
xx = 0.7
for h, w in cols:
    fill = PRIMARY_SOFT if h == "创想∞" else SURFACE
    c = PRIMARY if h == "创想∞" else FG
    card(s, xx, 2.05, w, 0.55, fill=fill)
    text(s, xx, 2.15, w, 0.4, [[(h, 13, c, True)]], align=PP_ALIGN.CENTER)
    xx += w + 0.05
# 表行
rows = [
    ("交互", "提问→回答", "想法→扩写", "多专家→红队→决策"),
    ("立场", "顺着用户", "顺着用户", "主动反驳、找失败理由"),
    ("输出", "单轮回答", "一份文案", "决策链+12份报告+诊断"),
    ("迭代", "无", "无", "二次审议、断点恢复"),
]
for r, row in enumerate(rows):
    xx = 0.7
    y = 2.65 + r * 0.78
    for j, (w, val) in enumerate(zip([c[1] for c in cols], row)):
        fill = PRIMARY_SOFT if j == 3 else SURFACE
        c = PRIMARY if j == 3 else (FG if j == 0 else MUTED)
        bold = j == 0 or j == 3
        card(s, xx, y, w, 0.68, fill=fill)
        text(s, xx + 0.1, y + 0.16, w - 0.2, 0.4, [[(val, 12, c, bold)]], align=PP_ALIGN.CENTER)
        xx += w + 0.05
text(s, 0.7, 6.1, 12, 0.6,
     [[("差异化来自：12 Agent 专业分工 + 红队对抗 + DAG 协同 + 工程化运行机制。", 14, PRIMARY, True)]])
footer(s, 12)

# ═══════════════════════════════════════════════
# P13 工程可靠性
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "RELIABILITY")
title(s, "不只是“能调用 Agent”，而是可控的工程系统")
grid = [
    ("Agent 依赖管理", PRIMARY), ("DAG 调度", PRIMARY), ("状态管理", PRIMARY),
    ("异常隔离", RED), ("Validator 校验", EXPERT), ("Critic/Repair 返工", EXPERT),
    ("Checkpoint 检查点", DECISION), ("Resume 断点续跑", DECISION), ("Report 持久化", DECISION),
]
for i, (name, c) in enumerate(grid):
    x = 0.7 + (i % 3) * 4.12
    y = 2.1 + (i // 3) * 1.3
    card(s, x, y, 3.9, 1.05, fill=SURFACE)
    text(s, x + 0.25, y + 0.32, 3.5, 0.5, [[("✓  " + name, 15, c, True)]])
text(s, 0.7, 6.15, 12, 0.6,
     [[("工程观：LLM 是不受信任的执行者，程序负责最终约束。", 16, PRIMARY, True)]])
footer(s, 13)

# ═══════════════════════════════════════════════
# P14 项目价值与未来扩展
# ═══════════════════════════════════════════════
s = slide()
eyebrow(s, "VALUE & ROADMAP")
title(s, "先做好“早期评审”这一件事")
# 当前价值
values = [
    ("对学生", "在投入时间金钱前，提前经历严苛、诚实的委员会拷问"),
    ("对教师", "12 维结构化初筛与可追溯报告，降低人工评审负荷"),
    ("对高校", "双创课程的可复盘教学工具，每个想法留下完整证据链"),
]
for i, (h, b) in enumerate(values):
    x = 0.7 + i * 4.12
    card(s, x, 2.05, 3.9, 1.7, fill=SURFACE)
    text(s, x + 0.25, 2.25, 3.5, 0.5, [[(h, 17, PRIMARY, True)]])
    text(s, x + 0.25, 2.85, 3.5, 0.8, [[(b, 12, MUTED, False)]])
# 未来扩展（明确标注）
card(s, 0.7, 4.1, 12, 2.0, fill=DECISION_SOFT, line=DECISION_SOFT)
text(s, 1.0, 4.25, 12, 0.4, [[("后续规划 / 未来扩展（非已完成）", 13, DECISION, True)]])
future = ["更多行业创业模板", "创业数据知识库", "用户长期项目记忆",
          "多轮项目修订", "项目历史版本对比", "更完整的模拟投资人/评委体系"]
for i, f in enumerate(future):
    x = 1.0 + (i % 3) * 3.9
    y = 4.75 + (i // 3) * 0.55
    text(s, x, y, 3.7, 0.4, [[("· " + f, 12, FG, False)]])
text(s, 0.7, 6.3, 12, 0.5,
     [[("不把未来规划包装成当前已有功能。", 12, MUTED, False)]])
footer(s, 14)

# ═══════════════════════════════════════════════
# P15 收尾
# ═══════════════════════════════════════════════
s = slide()
card(s, 0, 0, 13.333, 0.16, fill=PRIMARY, line=PRIMARY, round_=False)
text(s, 0.9, 2.6, 11.5, 1.2,
     [[("在见投资人之前，", 36, FG, True)],
      [("先见你的 AI 委员会。", 36, PRIMARY, True)]], sp_after=4)
text(s, 0.9, 4.6, 11.5, 0.6, [[("创想∞ AI创业委员会", 22, FG, True)]])
tags = ["12 Agent", "红队对抗", "DAG 协同", "工程化运行", "创业诊断闭环"]
xx = 0.9
for t in tags:
    chip(s, t, xx, 5.6, 2.1, PRIMARY_SOFT, PRIMARY)
    xx += 2.3

out = HERE / "07_路演PPT_v2.pptx"
prs.save(str(out))
print("saved:", out, out.stat().st_size, "bytes,", len(list(prs.slides)), "slides")
